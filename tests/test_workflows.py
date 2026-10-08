from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.identity.models import Role
from apps.identity.permissions import visible_projects
from apps.projects.models import ProgressSnapshot, Project, Risk
from apps.projects.services import capture_baseline, project_health
from apps.work.models import Task, TaskAttachment

pytestmark = pytest.mark.django_db


@pytest.fixture
def setup_project(db):
    pm = User.objects.create_user("owner", password="valid-password")
    pm.pcc_profile.role = Role.PROJECT_MANAGER
    pm.pcc_profile.save()
    member = User.objects.create_user("member", password="valid-password")
    stranger = User.objects.create_user("stranger", password="valid-password")
    management = User.objects.create_user("management", password="valid-password")
    management.pcc_profile.role = Role.MANAGEMENT
    management.pcc_profile.save()
    p = Project.objects.create(
        name="Private project",
        code="PRIVATE",
        manager=pm,
        start_date=date(2026, 10, 1),
        target_end_date=date(2026, 11, 1),
    )
    p.members.add(member)
    task = Task.objects.create(
        project=p,
        title="Build",
        assignee=member,
        start_date=date(2026, 10, 2),
        due_date=date(2026, 10, 10),
    )
    return p, task, pm, member, stranger, management


@pytest.mark.parametrize(
    "page",
    [
        "dashboard",
        "project_list",
        "project_detail",
        "project_progress",
        "project_timeline",
        "project_tasks",
        "project_issues",
        "project_milestones",
        "project_notes",
        "project_activity",
        "project_members",
        "reports",
        "timeline_global",
        "note_list",
        "project_edit",
        "risk_create",
    ],
)
def test_project_pages_render(client, setup_project, page):
    project, task, pm, *_ = setup_project
    client.force_login(pm)
    args = (
        [project.pk]
        if page.startswith("project_") and page not in ["project_list"] or page == "risk_create"
        else []
    )
    response = client.get(reverse("projects:" + page, args=args))
    assert response.status_code == 200


@pytest.mark.parametrize("page", ["task_detail", "task_update"])
def test_task_pages_render(client, setup_project, page):
    p, task, pm, *_ = setup_project
    client.force_login(pm)
    assert client.get(reverse("work:" + page, args=[task.pk])).status_code == 200


def test_member_only_sees_membership_and_own_tasks(client, setup_project):
    p, task, pm, member, stranger, _ = setup_project
    assert list(visible_projects(member)) == [p]
    client.force_login(stranger)
    assert client.get(reverse("projects:project_detail", args=[p.pk])).status_code == 404
    assert client.get(reverse("work:task_detail", args=[task.pk])).status_code == 404
    assert b"Private project" not in client.get(reverse("projects:dashboard")).content
    assert client.get(reverse("projects:export_report"), {"project": p.pk}).status_code == 404


def test_member_kanban_permission_and_planning_fields(client, setup_project):
    p, task, pm, member, *_ = setup_project
    client.force_login(member)
    board = client.get(reverse("projects:project_tasks", args=[p.pk]))
    assert b'data-can-drag="true"' in board.content
    assert b'data-draggable="true"' in board.content
    form = client.get(reverse("work:task_update", args=[task.pk])).context["form"]
    assert "weight" not in form.fields
    result = client.post(
        reverse("work:task_update", args=[task.pk]),
        {"status": "in_progress", "progress": "35", "weight": "999"},
    )
    assert result.status_code == 302
    task.refresh_from_db()
    assert task.progress == 35 and task.weight == 1
    result = client.patch(
        reverse("work:task_dates", args=[task.pk]),
        "start=2026-10-01&end=2026-10-04",
        content_type="application/x-www-form-urlencoded",
    )
    assert result.status_code == 403


def test_gantt_validates_and_persists_dates(client, setup_project):
    p, task, pm, *_ = setup_project
    client.force_login(pm)
    url = reverse("work:task_dates", args=[task.pk])
    for body in [
        "start=no&end=bad",
        "start=2026-10-10&end=2026-10-01",
        "start=2026-02-30&end=2026-03-01",
        "",
    ]:
        assert (
            client.patch(url, body, content_type="application/x-www-form-urlencoded").status_code
            == 400
        )
    task.refresh_from_db()
    assert task.due_date == date(2026, 10, 10)
    assert (
        client.patch(
            url, "start=2026-10-03&end=2026-10-12", content_type="application/x-www-form-urlencoded"
        ).status_code
        == 200
    )
    task.refresh_from_db()
    assert task.start_date == date(2026, 10, 3) and task.due_date == date(2026, 10, 12)
    rendered = client.get(reverse("projects:project_timeline", args=[p.pk])).content
    assert b"Edit jadwal" in rendered


def test_blocked_requires_reason_and_dependencies_gate_start(client, setup_project):
    p, task, pm, member, *_ = setup_project
    previous = Task.objects.create(project=p, title="First")
    task.dependencies.add(previous)
    client.force_login(member)
    url = reverse("work:task_status", args=[task.pk])
    assert client.post(url, {"status": "in_progress"}).status_code == 400
    assert client.post(url, {"status": "blocked"}).status_code == 400
    assert (
        client.post(
            url, {"status": "blocked", "blocked_reason": "Waiting for environment"}
        ).status_code
        == 200
    )
    previous.status = "done"
    previous.save()
    assert client.post(url, {"status": "done"}).status_code == 200


def test_task_create_dependencies_and_cycle_validation(client, setup_project):
    p, task, pm, member, *_ = setup_project
    client.force_login(pm)
    url = reverse("work:task_create") + f"?project={p.pk}"
    data = {
        "title": "Second",
        "status": "todo",
        "priority": "medium",
        "progress": 0,
        "weight": 2,
        "order": 0,
        "dependencies": [task.pk],
        "assignee": member.pk,
    }
    assert client.post(url, data).status_code == 302
    second = p.tasks.get(title="Second")
    assert list(second.dependencies.all()) == [task]
    from apps.work.forms import TaskForm

    data.update(title=task.title, dependencies=[second.pk])
    form = TaskForm(data, instance=task, project=p)
    assert not form.is_valid()
    assert "dependencies" in form.errors


def test_capture_baseline_and_daily_snapshots(setup_project):
    p, task, pm, *_ = setup_project
    capture_baseline(p, pm)
    task.refresh_from_db()
    assert task.baseline_end_date == date(2026, 10, 10)
    assert p.baseline_points.first().planned_percent == 0
    assert p.baseline_points.last().planned_percent == 100
    task.due_date = date(2026, 10, 14)
    task.progress = 50
    task.save()
    assert task.schedule_deviation == 4
    from django.utils import timezone

    snapshot = p.snapshots.get(date=timezone.localdate())
    assert snapshot.actual_progress == Decimal("50")
    call_command("snapshot_progress")
    assert p.snapshots.filter(date=timezone.localdate()).count() == 1


def test_risks_separate_from_issues(setup_project):
    p, *_ = setup_project
    p.issues.create(title="Actual issue", severity="critical")
    Risk.objects.create(project=p, title="Potential delay", severity="medium")
    health = project_health(p)
    assert health["critical_issues"] == 1
    assert health["risk_count"] == 1 and health["risk_top_severity"] == "medium"


@pytest.mark.parametrize("kind", ["csv", "xlsx", "pdf"])
@pytest.mark.parametrize(
    "dataset", ["summary", "progress", "timeline", "tasks", "issues", "milestones"]
)
def test_exports(client, setup_project, kind, dataset):
    p, task, pm, *_ = setup_project
    client.force_login(pm)
    response = client.get(reverse("projects:export_report"), {"dataset": dataset, "format": kind})
    assert response.status_code == 200
    assert f'.{kind}"' in response["Content-Disposition"]
    if kind == "pdf":
        assert response.content.startswith(b"%PDF")
    if kind == "xlsx":
        sheet = load_workbook(BytesIO(response.content)).active
        assert sheet.max_row >= 1


def test_export_spreadsheet_formula_escape(client, setup_project):
    p, task, pm, *_ = setup_project
    p.name = '=HYPERLINK("bad")'
    p.save()
    client.force_login(pm)
    response = client.get(reverse("projects:export_report"), {"format": "xlsx"})
    cell = load_workbook(BytesIO(response.content)).active["B2"]
    assert cell.data_type == "s" and cell.value.startswith("'=")


def test_local_auth_requires_password_and_does_not_change_role(client, setup_project):
    p, task, pm, member, *_ = setup_project
    url = reverse("identity:login")
    assert client.get(url).status_code == 200
    assert (
        client.post(url, {"username": "member", "password": "wrong", "role": "pm"}).status_code
        == 200
    )
    assert (
        client.post(
            url,
            {
                "username": "member",
                "password": "valid-password",
                "next": "https://untrusted.example/",
                "role": "pm",
            },
        ).url
        == "/"
    )
    member.pcc_profile.refresh_from_db()
    assert member.pcc_profile.role == "member"


def test_management_readonly(client, setup_project):
    p, task, pm, member, stranger, management = setup_project
    client.force_login(management)
    assert client.get(reverse("projects:project_detail", args=[p.pk])).status_code == 200
    assert (
        client.post(reverse("work:task_status", args=[task.pk]), {"status": "done"}).status_code
        == 403
    )
    assert (
        client.post(reverse("work:issue_create"), {"project": p.pk, "title": "Write"}).status_code
        == 403
    )


def test_membership_management_and_settings(client, setup_project):
    p, task, pm, member, stranger, *_ = setup_project
    client.force_login(pm)
    assert (
        client.post(
            reverse("projects:project_members", args=[p.pk]), {"members": [stranger.pk]}
        ).status_code
        == 302
    )
    assert visible_projects(stranger).filter(pk=p.pk).exists()
    assert client.get(reverse("identity:settings")).status_code == 200
    assert client.get(reverse("identity:user_create")).status_code == 403


def test_attachments_access(client, setup_project, tmp_path):
    p, task, pm, member, stranger, *_ = setup_project
    with override_settings(MEDIA_ROOT=tmp_path):
        client.force_login(member)
        response = client.post(
            reverse("work:attachment_upload", args=[task.pk]),
            {"file": SimpleUploadedFile("notes.txt", b"project notes")},
        )
        assert response.status_code == 302
        attachment = TaskAttachment.objects.get(task=task)
        response = client.get(reverse("work:attachment_download", args=[attachment.pk]))
        assert b"".join(response.streaming_content) == b"project notes"
        client.force_login(stranger)
        assert (
            client.get(reverse("work:attachment_download", args=[attachment.pk])).status_code == 404
        )
        client.force_login(member)
        assert (
            client.post(reverse("work:attachment_delete", args=[attachment.pk])).status_code == 302
        )
        assert not TaskAttachment.objects.exists()


def test_project_delete_does_not_recreate_snapshots(setup_project):
    p, *_ = setup_project
    p.delete()
    assert not ProgressSnapshot.objects.exists()


@pytest.mark.parametrize(
    "params",
    [
        {"from": "2026-02-30"},
        {"to": "bad"},
        {"progress_min": "Infinity"},
        {"progress_max": "NaN"},
        {"progress_min": "90", "progress_max": "10"},
    ],
)
def test_invalid_filters_recover(client, setup_project, params):
    p, task, pm, *_ = setup_project
    client.force_login(pm)
    response = client.get(reverse("projects:project_list"), params)
    assert response.status_code == 200
    assert response.context["rows"]


def test_member_reports_issue_without_setting_resolution(client, setup_project):
    p, task, pm, member, *_ = setup_project
    client.force_login(member)
    response = client.post(
        reverse("work:issue_create") + f"?project={p.pk}",
        {
            "project": p.pk,
            "title": "Problem",
            "severity": "high",
            "priority": "high",
            "status": "resolved",
            "resolution": "Injected",
            "owner": pm.pk,
        },
    )
    assert response.status_code == 302
    issue = p.issues.get(title="Problem")
    assert issue.status == "open" and issue.owner == member and issue.resolution == ""


def test_duplicate_phase_is_form_error(client, setup_project):
    p, task, pm, *_ = setup_project
    p.phases.create(name="Existing", sequence=1)
    client.force_login(pm)
    response = client.post(
        reverse("projects:phase_create", args=[p.pk]),
        {"name": "Duplicate", "sequence": 1, "status": "not_started"},
    )
    assert response.status_code == 200
    assert response.context["form"].errors.get("sequence")


def test_weighted_baseline_and_override_snapshot(client, setup_project):
    p, task, pm, *_ = setup_project
    Task.objects.create(
        project=p,
        title="Second",
        weight=3,
        start_date=date(2026, 10, 8),
        due_date=date(2026, 10, 15),
    )
    capture_baseline(p, pm)
    # 25 + 3/8 * 75 = 53.125, rounded half-even to the shared 0.01 precision.
    assert p.baseline_points.get(date=date(2026, 10, 10)).planned_percent == Decimal("53.12")
    p.progress_override = 77
    p.save()
    from django.utils import timezone

    assert p.snapshots.get(date=timezone.localdate()).actual_progress == 77


def test_staff_provisions_local_account_and_password_change(client, setup_project):
    p, task, pm, *_ = setup_project
    pm.is_staff = True
    pm.save()
    client.force_login(pm)
    response = client.post(
        reverse("identity:user_create"),
        {
            "username": "new-pm",
            "password1": "safe-password",
            "password2": "safe-password",
            "role": "pm",
        },
    )
    assert response.status_code == 302
    user = User.objects.get(username="new-pm")
    assert user.check_password("safe-password") and user.pcc_profile.role == "pm"
    client.force_login(user)
    response = client.post(
        reverse("identity:settings"),
        {
            "action": "password",
            "old_password": "safe-password",
            "new_password1": "updated-password",
            "new_password2": "updated-password",
        },
    )
    assert response.status_code == 302
    user.refresh_from_db()
    assert user.check_password("updated-password")


def test_member_cannot_manipulate_foreign_notes_issues_or_baseline(client, setup_project):
    p, task, pm, member, stranger, *_ = setup_project
    issue = p.issues.create(title="Private issue")
    note = p.notes.create(title="Private note", content="Secret")
    client.force_login(stranger)
    for url in [
        reverse("work:issue_update", args=[issue.pk]),
        reverse("projects:note_edit", args=[p.pk, note.pk]),
        reverse("projects:baseline_add", args=[p.pk]),
    ]:
        assert client.get(url).status_code == 404
    assert b"Private issue" not in client.get(reverse("work:issue_list")).content
    assert b"Private note" not in client.get(reverse("projects:note_list")).content


@pytest.mark.parametrize("actor_index", [2, 3])
def test_daily_update_keeps_planning_fields_and_validates_blockers(
    client, setup_project, actor_index
):
    project, task, *_ = setup_project
    client.force_login(setup_project[actor_index])
    url = reverse("work:task_update", args=[task.pk]) + "?mode=update"
    form = client.get(url).context["form"]
    assert set(form.fields) == {"status", "progress", "blocked_reason", "blocked_dependency"}
    response = client.post(url, {"status": "blocked", "progress": 30})
    assert response.status_code == 200
    assert response.context["form"].errors
    response = client.post(
        url,
        {
            "status": "in_progress",
            "progress": 35,
            "title": "Unintended rename",
            "weight": 999,
        },
    )
    assert response.status_code == 302
    task.refresh_from_db()
    assert task.title == "Build" and task.weight == 1 and task.progress == 35


def test_schedule_views_and_blocked_work_are_separate(client, setup_project):
    project, task, pm, member, *_ = setup_project
    client.force_login(member)
    task.status = "blocked"
    task.blocked_reason = "Waiting for access"
    task.save()
    response = client.get(reverse("work:my_tasks"))
    assert task in response.context["blocked"]
    assert task not in response.context["next_up"]
    url = reverse("projects:project_timeline", args=[project.pk])
    response = client.get(url)
    assert b'data-fullscreen-panel="Timeline saat ini"' in response.content
    assert b"?mode=update" in response.content
    assert b"task-plan-fill" in response.content and b"task-actual-fill" in response.content
    assert b"Perubahan tanggal dari baseline" in response.content


def test_collaborator_can_update_shared_task_without_double_counting(client, setup_project):
    project, task, pm, member, stranger, management = setup_project
    task.collaborators.add(stranger)
    before = project.actual_progress
    client.force_login(stranger)
    assert project in visible_projects(stranger)
    assert task in client.get(reverse("work:my_tasks")).context["next_up"]
    url = reverse("work:task_update", args=[task.pk]) + "?mode=update"
    form = client.get(url).context["form"]
    assert "collaborators" not in form.fields
    assert client.post(url, {"status": "in_progress", "progress": 35}).status_code == 302
    task.refresh_from_db()
    assert task.progress == 35 and task.collaborators.filter(pk=stranger.pk).exists()
    assert project.tasks.count() == 1 and task.weight == 1
    assert before == 0


def test_two_assignee_form_rejects_duplicate_or_third_person(setup_project):
    from apps.work.forms import TaskForm

    project, task, pm, member, stranger, *_ = setup_project
    project.members.add(stranger)
    base = {"title": "Build", "assignee": member.pk, "status": "todo", "progress": 0}
    for collaborators in [[member.pk], [stranger.pk, pm.pk]]:
        form = TaskForm({**base, "collaborators": collaborators}, instance=task, project=project)
        assert not form.is_valid() and "collaborators" in form.errors
    form = TaskForm({**base, "collaborators": [stranger.pk]}, instance=task, project=project)
    assert form.is_valid(), form.errors
    form.save()
    assert task.collaborators.filter(pk=stranger.pk).exists()


def test_pm_can_resolve_issue_without_owner(client, setup_project):
    from apps.work.models import Issue

    project, task, pm, *_ = setup_project
    issue = Issue.objects.create(project=project, title="Access blocked", severity="high")
    client.force_login(pm)
    response = client.post(
        reverse("work:issue_update", args=[issue.pk]),
        {
            "project": project.pk,
            "title": issue.title,
            "status": "resolved",
            "resolution": "Confirmed fixed by external team",
            "severity": "high",
            "priority": "high",
        },
    )
    assert response.status_code == 302
    issue.refresh_from_db()
    assert issue.status == "resolved" and issue.resolved_at and issue.owner is None
    assert issue.resolution == "Confirmed fixed by external team"


def test_timeline_plan_is_inclusive_and_unknown_dates_are_explicit(setup_project, monkeypatch):
    from apps.projects.services import timeline_data

    project, task, *_ = setup_project
    monkeypatch.setattr("apps.projects.services.timezone.localdate", lambda: date(2026, 10, 10))
    row = timeline_data(project)["timeline_rows"][0]
    assert row["planned_percent"] == 100
    task.start_date = None
    task.due_date = None
    task.save()
    row = timeline_data(project)["timeline_rows"][0]
    assert row["planned_percent"] is None and row["undated"]


def test_presentation_access_and_pm_configuration(client, setup_project):
    import json

    from apps.projects.models import BoardColumn, PresentationConfig
    from apps.projects.presentation_views import resolve_sections

    project, task, pm, member, stranger, management = setup_project
    column = BoardColumn.objects.create(project=project, title="Feedback")
    url = reverse("projects:presentation", args=[project.pk])
    play = reverse("projects:presentation_play", args=[project.pk])
    rows = resolve_sections(project, [])
    rows[-1]["enabled"] = True
    rows.insert(1, rows.pop())
    data = {
        "period_start": "2026-10-11",
        "period_end": "2026-10-25",
        "sections": json.dumps(rows),
        "action": "play",
    }
    client.force_login(stranger)
    assert client.get(play).status_code == 404
    client.force_login(member)
    assert client.get(url).status_code == 200
    assert client.get(play).status_code == 200
    assert client.post(url, data).status_code == 403
    client.force_login(pm)
    response = client.post(url, data)
    assert response.status_code == 302 and response.url == play
    assert (
        PresentationConfig.objects.get(project=project).sections[1]["key"] == f"column:{column.pk}"
    )
    assert client.get(play).context["slides"][1]["kind"] == "column"
    # Cross-project, duplicate, disabled-all and invalid dates never overwrite saved configuration.
    for mutate in ("foreign", "duplicate", "disabled", "dates"):
        bad = [dict(row) for row in rows]
        if mutate == "foreign":
            bad[0]["key"] = "column:999999"
        elif mutate == "duplicate":
            bad[0]["key"] = bad[1]["key"]
        elif mutate == "disabled":
            for row in bad:
                row["enabled"] = False
        response = client.post(
            url,
            {
                **data,
                "sections": json.dumps(bad),
                "period_end": "2026-10-01" if mutate == "dates" else data["period_end"],
            },
        )
        assert response.status_code == 200 and response.context["form"].errors
        assert PresentationConfig.objects.get(project=project).sections == rows


def test_feedback_does_not_change_progress_and_scopes_links(client, setup_project):
    from apps.projects.models import BoardColumn, FeedbackCard

    project, task, pm, member, *_ = setup_project
    column = BoardColumn.objects.create(project=project, title="Feedback")
    before = project.actual_progress
    client.force_login(pm)
    url = reverse("projects:feedback_create", args=[project.pk, column.pk])
    response = client.post(
        url,
        {
            "title": "Customer feedback",
            "description": "Needs discussion",
            "status": "agreed",
            "owner": member.pk,
            "linked_task": task.pk,
        },
    )
    assert response.status_code == 302
    assert FeedbackCard.objects.get(column=column).linked_task == task
    assert project.actual_progress == before
    other = Project.objects.create(
        name="Other",
        code="OTHER",
        manager=pm,
        start_date=project.start_date,
        target_end_date=project.target_end_date,
    )
    foreign = Task.objects.create(project=other, title="Other task")
    response = client.post(
        url, {"title": "Invalid link", "status": "open", "linked_task": foreign.pk}
    )
    assert response.status_code == 200 and response.context["form"].errors
    client.force_login(member)
    assert client.post(url, {"title": "Not allowed"}).status_code == 403


def test_custom_column_move_preserves_status_progress_and_validates_scope(client, setup_project):
    from apps.projects.models import BoardColumn

    project, task, pm, member, stranger, *_ = setup_project
    column = BoardColumn.objects.create(project=project, title="For discussion")
    task.status, task.progress = "in_progress", Decimal("45")
    task.save()
    before = project.actual_progress
    url = reverse("projects:board_move", args=[project.pk, task.pk])
    client.force_login(member)
    assert client.post(url, {"column": column.pk}).status_code == 200
    task.refresh_from_db()
    assert (task.status, task.progress, task.board_column_id) == (
        "in_progress",
        Decimal("45"),
        column.pk,
    )
    assert project.actual_progress == before
    other = Project.objects.create(
        name="Other",
        code="OTHER",
        manager=pm,
        start_date=project.start_date,
        target_end_date=project.target_end_date,
    )
    foreign = BoardColumn.objects.create(project=other, title="Foreign")
    assert client.post(url, {"column": foreign.pk}).status_code == 404
    assert (
        client.post(
            reverse("work:task_status", args=[task.pk]),
            {"status": "review", "clear_column": "true"},
        ).status_code
        == 200
    )
    task.refresh_from_db()
    assert task.board_column_id is None and task.status == "review"
    client.force_login(stranger)
    assert client.post(url, {"column": column.pk}).status_code == 404


def test_presentation_period_pagination_and_snapshot(client, setup_project):
    from apps.projects.models import BoardColumn, FeedbackCard, PresentationConfig
    from apps.projects.presentation_views import presentation_slides, resolve_sections

    project, task, pm, member, *_ = setup_project
    column = BoardColumn.objects.create(project=project, title="Feedback")
    for n in range(9):
        FeedbackCard.objects.create(column=column, title=f"Feedback {n}")
    rows = resolve_sections(project, [])
    rows[-1]["enabled"] = True
    start, end = date(2026, 10, 11), date(2026, 10, 25)
    scheduled = Task.objects.create(
        project=project, title="Next scheduled", start_date=start, due_date=end
    )
    active = Task.objects.create(
        project=project, title="Carry forward", status="in_progress", progress=55
    )
    Task.objects.create(project=project, title="No date backlog")
    Task.objects.create(
        project=project, title="Completed", status="done", start_date=start, due_date=end
    )
    slides = presentation_slides(project, rows, start, end)
    pages = [s for s in slides if s["kind"] == "column"]
    assert [len(s["items"]) for s in pages] == [4, 4, 1]
    upcoming = next(s for s in slides if s["kind"] == "next")
    assert {t.pk for t in upcoming["items"]} == {scheduled.pk, active.pk}
    PresentationConfig.objects.create(
        project=project, sections=rows, period_start=start, period_end=end
    )
    client.force_login(pm)
    first = client.get(reverse("projects:presentation_play", args=[project.pk])).content
    FeedbackCard.objects.filter(column=column).update(title="Updated later")
    assert b"Feedback 0" in first and b"Updated later" not in first
    assert (
        b"Updated later"
        in client.get(reverse("projects:presentation_play", args=[project.pk])).content
    )


def test_presentation_long_prose_is_complete_and_malformed_source_is_rejected(
    client, setup_project
):
    import json

    from apps.projects.models import BoardColumn, FeedbackCard
    from apps.projects.presentation_views import (
        PresentationForm,
        presentation_slides,
        resolve_sections,
    )

    project, task, pm, *_ = setup_project
    column = BoardColumn.objects.create(project=project, title="Long feedback")
    prose = "A decision and its supporting context. " * 100
    FeedbackCard.objects.create(column=column, title="Detailed discussion", description=prose)
    rows = resolve_sections(project, [])
    rows[-1]["enabled"] = True
    slides = presentation_slides(project, rows, date(2026, 10, 11), date(2026, 10, 25))
    descriptions = [
        item["feedback"].description
        for slide in slides
        if slide["kind"] == "column"
        for item in slide["items"]
    ]
    assert len(descriptions) > 1 and "".join(descriptions) == prose
    assert FeedbackCard.objects.get(column=column).description == prose
    rows[0]["key"] = ["invalid"]
    form = PresentationForm(
        {"sections": json.dumps(rows), "period_start": "2026-10-11", "period_end": "2026-10-25"},
        project=project,
    )
    assert not form.is_valid() and "sections" in form.errors
    client.force_login(pm)
    assert (
        client.post(
            reverse("projects:board_move", args=[project.pk, task.pk]), {"column": "invalid"}
        ).status_code
        == 400
    )


def test_feedback_moves_only_within_project_and_by_pm(client, setup_project):
    from apps.projects.models import BoardColumn, FeedbackCard

    project, task, pm, member, *_ = setup_project
    source = BoardColumn.objects.create(project=project, title="Feedback")
    target = BoardColumn.objects.create(project=project, title="Agreed feedback")
    card = FeedbackCard.objects.create(column=source, title="Discuss", linked_task=task)
    url = reverse("projects:feedback_move", args=[project.pk, card.pk])
    client.force_login(member)
    assert client.post(url, {"column": target.pk}).status_code == 403
    client.force_login(pm)
    assert client.post(url, {"column": target.pk}).status_code == 200
    card.refresh_from_db()
    assert card.column == target and card.linked_task == task
    other = Project.objects.create(
        name="Other",
        code="OTHER",
        manager=pm,
        start_date=project.start_date,
        target_end_date=project.target_end_date,
    )
    foreign = BoardColumn.objects.create(project=other, title="Other feedback")
    assert client.post(url, {"column": foreign.pk}).status_code == 404
    assert client.post(url, {"column": "invalid"}).status_code == 400
    card.refresh_from_db()
    assert card.column == target


def test_status_update_keeps_custom_group_and_blocked_drawer_can_move_to_status_column(
    client, setup_project
):
    from apps.projects.models import BoardColumn

    project, task, pm, member, *_ = setup_project
    column = BoardColumn.objects.create(project=project, title="Discuss")
    task.board_column = column
    task.save()
    client.force_login(member)
    assert (
        client.post(reverse("work:task_status", args=[task.pk]), {"status": "review"}).status_code
        == 200
    )
    task.refresh_from_db()
    assert task.board_column_id == column.pk
    url = (
        reverse("work:task_update", args=[task.pk])
        + "?mode=update&board_status=blocked&clear_column=true"
    )
    response = client.get(url)
    assert response.context["form"]["status"].value() == "blocked"
    response = client.post(
        url, {"status": "blocked", "progress": "25", "blocked_reason": "Waiting for infrastructure"}
    )
    assert response.status_code == 302
    task.refresh_from_db()
    assert task.status == "blocked" and task.board_column_id is None
    assert task.blocked_reason == "Waiting for infrastructure"
