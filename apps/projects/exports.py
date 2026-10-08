"""Permission-scoped project exports, shared across CSV, Excel, and PDF."""

import csv
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone

from apps.identity.permissions import visible_projects

from .services import actual_progress, planned_progress_on, variance


def export_rows(projects, dataset):
    if dataset == "summary":
        headers = [
            "Code",
            "Project",
            "PM",
            "Status",
            "Phase",
            "Start",
            "Target",
            "Planned %",
            "Actual %",
            "Variance pp",
            "Open issues",
            "Open risks",
            "Next milestone",
            "Current work",
            "Next work",
        ]
        rows = []
        for p in projects:
            phase, milestone = p.effective_current_phase, p.next_milestone
            rows.append(
                [
                    p.code,
                    p.name,
                    str(p.manager or ""),
                    p.get_status_display(),
                    str(phase or ""),
                    p.start_date,
                    p.target_end_date,
                    planned_progress_on(p),
                    actual_progress(p),
                    variance(p),
                    p.issues.filter(status__in=["open", "in_progress", "blocked"]).count(),
                    p.risks.filter(status="open").count(),
                    str(milestone or ""),
                    "; ".join(p.tasks.filter(status="in_progress").values_list("title", flat=True)),
                    "; ".join(
                        p.tasks.filter(status__in=["todo", "backlog"]).values_list(
                            "title", flat=True
                        )
                    ),
                ]
            )
        return headers, rows
    if dataset == "progress":
        return ["Code", "Date", "Planned %", "Actual %", "Variance pp"], [
            [p.code, s.date, s.planned_progress, s.actual_progress, s.variance]
            for p in projects
            for s in p.snapshots.all()
        ]
    if dataset in ("tasks", "timeline"):
        return [
            "Code",
            "Task",
            "Phase",
            "Assignee",
            "Status",
            "Priority",
            "Start",
            "Due",
            "Progress %",
            "Weight",
            "Baseline start",
            "Baseline end",
            "Deviation days",
            "Dependencies",
            "Blocked reason",
        ], [
            [
                p.code,
                t.title,
                str(t.phase or ""),
                t.assignee_names,
                t.get_status_display(),
                t.get_priority_display(),
                t.start_date,
                t.due_date,
                t.display_progress,
                t.weight,
                t.baseline_start_date,
                t.baseline_end_date,
                t.schedule_deviation,
                "; ".join(t.dependencies.values_list("title", flat=True)),
                t.blocked_reason,
            ]
            for p in projects
            for t in p.tasks.select_related("phase", "assignee").prefetch_related(
                "dependencies", "collaborators"
            )
        ]
    if dataset == "issues":
        return ["Code", "Issue", "Severity", "Priority", "Status", "Owner", "Due", "Resolution"], [
            [
                p.code,
                i.title,
                i.get_severity_display(),
                i.get_priority_display(),
                i.get_status_display(),
                str(i.owner or ""),
                i.due_date,
                i.resolution,
            ]
            for p in projects
            for i in p.issues.select_related("owner")
        ]
    if dataset == "milestones":
        return ["Code", "Milestone", "Date", "Status", "Owner", "Description"], [
            [p.code, m.name, m.date, m.get_status_display(), str(m.owner or ""), m.description]
            for p in projects
            for m in p.milestones.select_related("owner")
        ]
    raise Http404("Jenis report tidak tersedia.")


def safe_cell(value):
    if value is None:
        return ""
    if not isinstance(value, str):
        return str(value)
    # Never allow untrusted titles to become spreadsheet formulas.
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


@login_required
def export_report(request):
    projects = (
        visible_projects(request.user)
        .select_related("manager")
        .prefetch_related("baseline_points", "tasks", "phases", "milestones")
    )
    if request.GET.get("project"):
        try:
            project_id = int(request.GET["project"])
        except ValueError as exc:
            raise Http404("Project tidak valid.") from exc
        project = get_object_or_404(projects, pk=project_id)
        projects = projects.filter(pk=project.pk)
    dataset = request.GET.get("dataset", "summary")
    headers, rows = export_rows(projects, dataset)
    kind = request.GET.get("format", "csv")
    filename = f"pcc-{dataset}-{timezone.localdate().isoformat()}"
    if kind == "csv":
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows([[safe_cell(v) for v in row] for row in rows])
        response = HttpResponse(
            "\ufeff" + output.getvalue(), content_type="text/csv; charset=utf-8"
        )
    elif kind == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = dataset.title()
        sheet.append(headers)
        for row in rows:
            sheet.append([safe_cell(v) if isinstance(v, str) else v for v in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0075DE")
        for i, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(i)].width = max(16, min(36, len(header) + 8))
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    elif kind == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A3, A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import (
            Image,
            LongTable,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            TableStyle,
        )

        output = BytesIO()
        size = landscape(A3 if len(headers) > 10 else A4)
        doc = SimpleDocTemplate(
            output, pagesize=size, leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24
        )
        styles = getSampleStyleSheet()
        cell_style = ParagraphStyle("Cell", fontSize=7, leading=9, wordWrap="CJK")
        table_data = [
            [Paragraph(escape(safe_cell(v)), cell_style) for v in row] for row in [headers, *rows]
        ]
        table = LongTable(
            table_data, colWidths=[(size[0] - 48) / len(headers)] * len(headers), repeatRows=1
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6f3fe")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#0075de")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#f6f5f4")],
                    ),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        doc.build(
            [
                Image(
                    str(settings.BASE_DIR / "static/brand/scope-horizontal.png"),
                    width=140,
                    height=45.37,
                    hAlign="LEFT",
                ),
                Spacer(1, 10),
                Paragraph(f"Scope — {dataset.title()}", styles["Title"]),
                Paragraph(str(timezone.localdate()), styles["Normal"]),
                Spacer(1, 16),
                table,
            ]
        )
        response = HttpResponse(output.getvalue(), content_type="application/pdf")
    else:
        raise Http404("Format export tidak tersedia.")
    response["Content-Disposition"] = f'attachment; filename="{filename}.{kind}"'
    return response
